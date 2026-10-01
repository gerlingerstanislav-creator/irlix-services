from __future__ import annotations

import base64
import os
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .schema import CanonicalCv


ACCENT = "008E8C"
TEXT = "2A3137"
MUTED = "A5A5A5"
LOGO_B64 = "iVBORw0KGgoAAAANSUhEUgAAAJ0AAABNCAYAAAC8GGb5AAAACXBIWXMAAAsTAAALEwEAmpwYAAAAAXNSR0IArs4c6QAAAARnQU1BAACxjwv8YQUAAAnBSURBVHgB7Z27ctxGFob/nmHtJVputKHBaFcbUU8gKNhaZ6ZrrcARwScgubJU5YhQaIs0qScgmLlsV4nOrIhQ5sx02VW2I8OZHWmc6UIO3Ge6YUIYNHAawIwHnP6qWpwBDi4anPm7+5zuHgGHowVvfvfYSzE+A1KPYR4/ufHO7QEcjhaMMd5kOhwEXm2pvw5HQ5TKXf7IsxbRkxv/mzidUzpHY8bpq5BrK/DywdVrh6MBTVWOcErnaARf5dIkr3KEczqHNaRyQgw2ObayKj354sa7SX6bczqHNbJaPWZaJl/cuBMWtzqnc1jxn28/CeQfn2MrBiIs2+6czmHFYDjc41lKlfvnOyel54DDwUSpHDMQbFA5YkWWEPVEsiSGfauy7KBbnsnyk77mObolkMWrsYl1mRcB2t7T/v4GRLoODnfvhWiAUrmUYUkqd+fEtJecjiOXMaqdjim5jRjJcirLCbpxBOp1+Qy7GPOj/T2l6QgDwXsOBw9H0vGOYMF/v/9sW17D49jKHutW1f4+VK/k1IEsMqkMCkb6cExz714sRShm2QrpnIeHq2BCIRKkglmbiUj2WOMqi7616Two55ulsvaXcfqAabmK8QW7SZSK8TY/qf+y9h762pEIZWHGipYIO7Xb5qidUrnUQuVeDwSX0efeawDneNMMXm0xLVlq1zSpX0XfQyYBuu8595vd9xP57yOWrVI7z7TbJt0lFfYRR+WI6xCno/Ydu1G8FIhhCNXrr4PUztg+tkrqi1fs3vB1cLpZxAn7ze7uSLbDuGoX4OFDv7j5zW8+9WehcsQK5gcFeTnfPoKCnDbqJXtXOLI4//VnsHKE9JI+l/rPUcX34vym5xcr538aXo6EqDs+TZ78+45VzG+eSve2LLeZ5e9QAUauE9EHswHHFVZqJ2OfBbWLb749EqgPwVSlu0wscvUayXITduroyDNJd6UJy7Ykm6EUrOp4c1K/8lJYbBLUpFRy3IJjmpQZMCa1+2h/c3qz+fNvonJEHzoSlHflqJ0HxzR370fsgDHSsLhFp7TiEuO4icoRfem9cv5zq3Chk3LY6THh4eBhOLW1pG2XzWFtQl+czqZD4SjSMj02rXa8dJcJN4hzWWg5GCCvbNx0l4m+OJ3HtHNxOhOkdkh5bbBStXs3SVNq5rRTOaIvTsftmTqnq0JchGCnx8ZTIZSXL4Y7bVWO6IPTBeApXdfD2q8fNBiAHTBOd4qDAShg3FbliEV3Ogr4HjJtndNxoPQYt0YYX85k6NiiOl027+IM/B7pKRz1tEyPdcE8E/7cb40H+0BvIsvncPBoORig9eUxP3xm8WCPczgbSO14AfeZqN11iNMlUMOaHFakb7FNh4LbrmZxHZwuhHlOrqOMgw+DScqLz3rZYICm9N3pKGbUKOm8tFAYRDAnZb9GGtrMla2iz05HDhfCYUd6uWmpchp5jMVc2Sr66HTUCN6Fczh7VLA3RFOYc2Xr6JvTJbKswXUcmnE5SYO1YbULtevjshJuWHoTSOUGon1noGauLId5BofpG/Jrxf43wJP+zoOVSwGp3ECgA7K5so0Hcc7T6SiAm9TY+KhflSmzieHgcfjhOlLRWchDz5U9UcOl7Fm06pU7bMat2mTDeMAM7k5mfvEGAwxE42ewaE4Xg5e496EWEnTUQYFgUVt7aEQ4j8EAi9iR2GXahXDUww4ES5X7/3snauhT87myrMOweCTgZRk8OLWrxirdpeewqqFP/LmyDdRunh0JG0LwHIriddRB6XqYOiXD30B3xPgj0nX8dNf5ROUyaK7sgcy1cqrl4WTI2hosWFSnS6AeUp3jZSs2heiWdXQfD5yv0320Tw7nsWwvxtPhD5o9NhQ+ahE0V3bHZuHsRQ4Oh0w73mDEZWISvE0Dlq1II9y/Pz3Uf4YLZy+y0yXgz+zvJBF9bZiku4THM74wt99mtHD2oqfBQvDaa07tMmzSXaRyarnYcmawcDbRh1WbOHEjp3YZVkn9i3ol63jh7Mkpsfhwp8w5tbNROQqLVKlcht3C2XucwQB9cDpyOKd2HNgqJ4O/g4sIXPgLZ9Nc2dqUW1+GNnHVjh8muG4cyiAte+iSDARzVC7Dbq7sRl3AmOJ0EeoZ1eyL0O4cnGMpPcZZ04Tia0nFfuoRP8V8Oa/ZzxmBU32OMQVyZceAw9179jHDDufKdjLAyrEkUBBYMKcjXqa3TUOf3Pp0Dj6TrEP7wQDO6Rx2jAVvFJBh4Wy1y+Gw5WD/jDdGT6qiWLmpl7H4Had0DntsFs4uXUq2X3hQvdNsdHFQYXuO6h5f9is7HtQvZlOPNjHY+ro8Q3lPc12f59RwvKePz36eILt2xkjfa1JybNW5NwzHZeePcq/jEjuvcF982GonzyuGa3m16+MUxHzvaTNXaPte7v16zXm+0vY0bi6Acrwdg91jbbeh7YqN5A1t46OcEGqptCzcsKrfZ/e6ra9T1jM8Lhybh+zPSvZl5ydG+r6ODfd1C03CWS0GA/S9er2dK/TBPci9jyqOowdFsSoafLiVO4Yeol+w+7pgR6/pQyxrJG+XbPNgHheYv3/6SaqgcP1M5c5hVnXaXzdYc0efy89tC6Acjus8r9Ni4exlbNMF+m9Y2B5Dpds8/d6HUgyyyytBAvWgiqr4NHdMHh+86ZIJVNXt57bRNU71dtPSXnQvAap/kC/7QuYVkxz1EdqseNVw4exldDp6eKZvKD3kSL/2YW7nkQ0pR97BqNo9x7Qz7oGXsSGoCs8/xFv6XiNcqV6RBEqFj1GdAjzStju4UsYjtKHhwtnL6HTkKNw2zDPD9lHuXHlITfJVrK//mpzcyxWq2vOdpOzYGFcdjcBwnlN9jWNUk91fCFWlt6fBwtmLOkdilvyE9kOgTMfH+q+vXweoHv18lnudQDlCot8HUA8z0O/J6TKHKYO2U2ckq5JN90dFoKuFJKlXehiu4fmf2Z/pMjodPTxT+yiAehixLtQj3Sqx29B2Sck+qm729D7qQKzBjGnfqj72HNOdEB/lbcSRvtfHqA4V0XoyKbpkNxzBoge8jNVrBFWNFR+mB+Usf9PvY6gPstgzXNXbTApG1Y0PpTwRmilK5tTUo833cE9Q3WGItc1jLDDLqHTZMKkIyjmy6ackFQqv1I7/QR8Fn97itcVKKw4f6ztmrabsnspcqrvJ4RZWWjfW1jgUdRD9A+qHr4s2S709p9RDzkNKQIpyT9keQHVyI4KdnSuzAn/JctfZPlAlvdLzknO+71+TbG9Xwrno/uLZXmu37+AOZRC1e7HmHasRJa/6vv6GeX/Zzr/U3392HD+H3L36nA4HA6Ho//8Bqu0KW551DE0AAAAAElFTkSuQmCC"


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _run(paragraph, text: str, *, size=10, bold=False, color=TEXT, font="Arial"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _heading(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=17, bold=True, color=TEXT)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), ACCENT)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def _bullet(cell_or_doc, text: str, size=10):
    p = cell_or_doc.add_paragraph(style=None) if hasattr(cell_or_doc, "add_paragraph") else None
    if p is None:
        return
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.first_line_indent = Cm(-0.15)
    p.paragraph_format.space_after = Pt(2)
    _run(p, f"• {text}", size=size)


def _prepare_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    return doc


def render_docx(cv: CanonicalCv, target: Path) -> None:
    doc = _prepare_document()

    logo_path = Path(tempfile.gettempdir()) / "irlix-cv-logo.png"
    if not logo_path.exists():
        logo_path.write_bytes(base64.b64decode(LOGO_B64))

    top = doc.add_table(rows=1, cols=2)
    top.autofit = False
    top.columns[0].width = Cm(4.2)
    top.columns[1].width = Cm(13.7)
    c0, c1 = top.rows[0].cells
    c0.width = Cm(4.2)
    c1.width = Cm(13.7)
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p_logo = c0.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_logo.add_run().add_picture(str(logo_path), width=Inches(1.25))
    p_name = c1.paragraphs[0]
    _run(p_name, cv.full_name or "CV", size=34, bold=True)
    if cv.position:
        p = c1.add_paragraph()
        _run(p, cv.position, size=17, color=ACCENT, bold=True)

    if cv.commercial_experience or cv.role_experience:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        if cv.commercial_experience:
            _run(p, f"Опыт коммерческой разработки: {cv.commercial_experience}", size=10.5, bold=True)
        if cv.role_experience:
            p2 = doc.add_paragraph()
            _run(p2, cv.role_experience, size=10.5, bold=True)

    if cv.summary:
        _heading(doc, "Профиль")
        for line in cv.summary:
            _bullet(doc, line)

    if cv.skill_groups or cv.tools:
        _heading(doc, "Навыки и умения")
        skill_table = doc.add_table(rows=1, cols=2)
        skill_table.autofit = False
        left, right = skill_table.rows[0].cells
        groups = cv.skill_groups
        split = (len(groups) + 1) // 2
        for cell, subset in ((left, groups[:split]), (right, groups[split:])):
            _set_cell_margins(cell, 40, 80, 40, 120)
            for group in subset:
                p = cell.add_paragraph()
                _run(p, group.title, size=10.5, bold=True, color=ACCENT)
                for item in group.items:
                    _bullet(cell, item)
        if cv.tools:
            p = right.add_paragraph()
            _run(p, "Используемые инструменты", size=10.5, bold=True, color=ACCENT)
            for item in cv.tools:
                _bullet(right, item)

    if cv.education or cv.languages or cv.certifications:
        _heading(doc, "Образование")
        info = doc.add_table(rows=1, cols=2)
        left, right = info.rows[0].cells
        for item in cv.education:
            p = left.add_paragraph()
            _run(p, item.institution or "Образование", bold=True)
            details = ", ".join([v for v in [item.degree, item.specialty, item.period, item.status] if v])
            if details:
                p2 = left.add_paragraph()
                _run(p2, details)
        if cv.languages:
            p = right.add_paragraph()
            _run(p, "Иностранные языки", bold=True, color=ACCENT)
            for item in cv.languages:
                _bullet(right, f"{item.language}{' — ' + item.level if item.level else ''}")
        if cv.certifications:
            p = right.add_paragraph()
            _run(p, "Сертификаты", bold=True, color=ACCENT)
            for item in cv.certifications:
                _bullet(right, item)

    if cv.projects:
        doc.add_section(WD_SECTION.NEW_PAGE)
        _heading(doc, f"Опыт работы{(' — ' + cv.commercial_experience) if cv.commercial_experience else ''}")
        for idx, project in enumerate(cv.projects):
            table = doc.add_table(rows=1, cols=2)
            table.autofit = False
            left, right = table.rows[0].cells
            left.width = Cm(5.2)
            right.width = Cm(12.7)
            _set_cell_shading(left, "F5F7F8")
            _set_cell_margins(left, 120, 120, 120, 160)
            _set_cell_margins(right, 120, 180, 120, 80)
            p = left.paragraphs[0]
            _run(p, project.role or "Роль", size=10.5, bold=True, color=ACCENT)
            if project.period:
                p = left.add_paragraph(); _run(p, project.period, bold=True)
            if project.team:
                p = left.add_paragraph(); _run(p, "Команда", size=9, bold=True, color=MUTED)
                p = left.add_paragraph(); _run(p, project.team, size=9.5)
            if project.technologies:
                p = left.add_paragraph(); _run(p, "Технологии", size=9, bold=True, color=MUTED)
                p = left.add_paragraph(); _run(p, ", ".join(project.technologies), size=9.5)

            p = right.paragraphs[0]
            _run(p, project.project_name or "Проект", size=10.5, bold=True)
            if project.description:
                p = right.add_paragraph(); _run(p, "Описание проекта", size=9.5, bold=True, color=ACCENT)
                p = right.add_paragraph(); _run(p, project.description, size=10)
            if project.responsibilities:
                p = right.add_paragraph(); _run(p, "Выполняемые задачи", size=9.5, bold=True, color=ACCENT)
                for item in project.responsibilities:
                    _bullet(right, item)
            if idx != len(cv.projects) - 1:
                doc.add_paragraph().paragraph_format.space_after = Pt(4)

    if cv.additional_info:
        doc.add_section(WD_SECTION.NEW_PAGE)
        _heading(doc, "Дополнительная информация")
        for item in cv.additional_info:
            p = doc.add_paragraph(); _run(p, item.title, bold=True, color=ACCENT)
            p = doc.add_paragraph(); _run(p, item.text)

    doc.save(target)


def render_pdf(docx_path: Path, target_dir: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["HOME"] = "/tmp"
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(target_dir), str(docx_path)],
        capture_output=True,
        text=True,
        timeout=120,
        env=env,
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice PDF conversion failed: {result.stderr.strip()}")
    pdf_path = target_dir / f"{docx_path.stem}.pdf"
    if not pdf_path.exists():
        raise RuntimeError("LibreOffice did not create PDF")
    return pdf_path
