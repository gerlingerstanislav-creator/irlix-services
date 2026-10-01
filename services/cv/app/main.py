from __future__ import annotations

import io
import os
import re
import shutil
import time
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pypdf import PdfReader

from .providers import ProviderResult, get_provider
from .render import render_docx, render_pdf
from .schema import CanonicalCv, ConvertRequest


APP_ROOT = "/api"
RENDER_DIR = Path(os.getenv("CV_RENDER_DIR", "/tmp/cv-converter"))
RENDER_TTL_SECONDS = int(os.getenv("CV_RENDER_TTL_SECONDS", "3600"))
MAX_CHUNK_CHARS = int(os.getenv("CV_LLM_CHUNK_CHARS", "6500"))
MAX_UPLOAD_BYTES = int(os.getenv("CV_MAX_UPLOAD_BYTES", str(20 * 1024 * 1024)))

app = FastAPI(title="IRLIX CV Converter", version="0.2.0")


def _split_text(text: str) -> list[str]:
    clean = text.replace("\r\n", "\n").strip()
    if len(clean) <= MAX_CHUNK_CHARS:
        return [clean]
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", clean) if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        parts = [paragraph[i : i + MAX_CHUNK_CHARS] for i in range(0, len(paragraph), MAX_CHUNK_CHARS)] if len(paragraph) > MAX_CHUNK_CHARS else [paragraph]
        for part in parts:
            candidate = f"{current}\n\n{part}".strip() if current else part
            if current and len(candidate) > MAX_CHUNK_CHARS:
                chunks.append(current)
                current = part
            else:
                current = candidate
    if current:
        chunks.append(current)
    return chunks


def _dedupe_strings(items: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for item in items:
        normalized = re.sub(r"\s+", " ", item or "").strip()
        key = normalized.casefold()
        if normalized and key not in seen:
            seen.add(key)
            result.append(normalized)
    return result


def _merge_cv(parts: list[CanonicalCv]) -> CanonicalCv:
    base = CanonicalCv()
    for part in parts:
        for field in ("full_name", "position", "commercial_experience", "role_experience"):
            if not getattr(base, field) and getattr(part, field):
                setattr(base, field, getattr(part, field))
        base.summary = _dedupe_strings(base.summary + part.summary)
        base.tools = _dedupe_strings(base.tools + part.tools)
        base.certifications = _dedupe_strings(base.certifications + part.certifications)

        existing_skill_groups = {g.title.strip().casefold(): g for g in base.skill_groups}
        for group in part.skill_groups:
            key = group.title.strip().casefold()
            if not key:
                continue
            if key in existing_skill_groups:
                existing_skill_groups[key].items = _dedupe_strings(existing_skill_groups[key].items + group.items)
            else:
                base.skill_groups.append(group)
                existing_skill_groups[key] = group

        def add_unique(target, incoming, key_fn):
            seen = {key_fn(item) for item in target}
            for item in incoming:
                key = key_fn(item)
                if key not in seen:
                    target.append(item)
                    seen.add(key)

        add_unique(base.education, part.education, lambda x: (x.institution + "|" + (x.specialty or "")).casefold())
        add_unique(base.languages, part.languages, lambda x: x.language.casefold())
        add_unique(base.projects, part.projects, lambda x: (x.project_name + "|" + (x.period or "") + "|" + x.role).casefold())
        add_unique(base.additional_info, part.additional_info, lambda x: (x.title + "|" + x.text[:60]).casefold())

        for field in ("email", "phone", "telegram", "location"):
            if not getattr(base.contacts, field) and getattr(part.contacts, field):
                setattr(base.contacts, field, getattr(part.contacts, field))
        base.contacts.links = _dedupe_strings(base.contacts.links + part.contacts.links)
    return base


def _cleanup_old_renders() -> None:
    RENDER_DIR.mkdir(parents=True, exist_ok=True)
    cutoff = time.time() - RENDER_TTL_SECONDS
    for child in RENDER_DIR.iterdir():
        try:
            if child.stat().st_mtime < cutoff:
                if child.is_dir():
                    shutil.rmtree(child, ignore_errors=True)
                else:
                    child.unlink(missing_ok=True)
        except FileNotFoundError:
            pass


def _extract_docx(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        blocks.append(text)
    return "\n\n".join(blocks)


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _extract_upload(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".doc"):
        raise HTTPException(status_code=415, detail="Формат .doc пока не поддерживается. Сохраните файл как .docx.")
    if lower.endswith(".docx"):
        return _extract_docx(data)
    if lower.endswith(".pdf"):
        text = _extract_pdf(data)
        if not text.strip():
            raise HTTPException(status_code=422, detail="В PDF не найден текстовый слой. Для сканов потребуется OCR.")
        return text
    raise HTTPException(status_code=415, detail="Поддерживаются файлы PDF и DOCX.")


async def _extract_canonical(text: str) -> tuple[CanonicalCv, dict]:
    chunks = _split_text(text)
    provider = get_provider()
    extracted: list[CanonicalCv] = []
    total_ms = 0
    model_name = provider.model
    provider_name = provider.name
    for index, chunk in enumerate(chunks, start=1):
        decorated = f"Часть CV {index} из {len(chunks)}. Извлеки только факты из этой части.\n\n{chunk}"
        result: ProviderResult = await provider.extract(decorated)
        extracted.append(result.cv)
        total_ms += result.inference_ms
        model_name = result.model
        provider_name = result.provider
    return _merge_cv(extracted), {
        "provider": provider_name,
        "model": model_name,
        "chunks": len(chunks),
        "input_chars": len(text),
        "inference_ms": total_ms,
        "total_ms": total_ms,
    }


@app.get(f"{APP_ROOT}/health")
async def health():
    provider = get_provider()
    return {"status": "ok", "provider": provider.name, "model": provider.model, "chunk_chars": MAX_CHUNK_CHARS}


@app.post(f"{APP_ROOT}/parse")
async def parse(file: UploadFile = File(...)):
    started = time.monotonic()
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Файл слишком большой.")
    text = _extract_upload(file.filename or "cv", data)
    if not text.strip():
        raise HTTPException(status_code=422, detail="В документе не найден текст.")
    try:
        cv, metrics = await _extract_canonical(text)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Не удалось разобрать CV локальной моделью: {exc}") from exc
    metrics["total_ms"] = int((time.monotonic() - started) * 1000)
    return {"cv": cv.model_dump(), "metrics": metrics}


@app.post(f"{APP_ROOT}/convert")
async def convert(request: ConvertRequest):
    try:
        cv, metrics = await _extract_canonical(request.text)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"LLM extraction failed: {exc}") from exc
    return {"cv": cv.model_dump(), "metrics": metrics}


def _safe_name(cv: CanonicalCv, suffix: str) -> str:
    base = re.sub(r"[^\w.-]+", "_", cv.full_name or "CV", flags=re.UNICODE).strip("_") or "CV"
    return f"{base}_IRLIX.{suffix}"


def _render_target(cv: CanonicalCv) -> tuple[Path, Path]:
    _cleanup_old_renders()
    folder = RENDER_DIR / uuid.uuid4().hex
    folder.mkdir(parents=True, exist_ok=True)
    docx_path = folder / _safe_name(cv, "docx")
    render_docx(cv, docx_path)
    return folder, docx_path


@app.post(f"{APP_ROOT}/render/docx")
async def render_docx_endpoint(cv: CanonicalCv):
    try:
        _, docx_path = _render_target(cv)
        return FileResponse(docx_path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=docx_path.name)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Не удалось сформировать DOCX: {exc}") from exc


@app.post(f"{APP_ROOT}/render/pdf")
async def render_pdf_endpoint(cv: CanonicalCv):
    try:
        folder, docx_path = _render_target(cv)
        pdf_path = render_pdf(docx_path, folder)
        return FileResponse(pdf_path, media_type="application/pdf", filename=pdf_path.name)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Не удалось сформировать PDF: {exc}") from exc
