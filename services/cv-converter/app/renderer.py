from __future__ import annotations

import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .models import CanonicalCv, ProjectItem

ACCENT = '008E8C'
TEXT = RGBColor(35, 35, 42)
MUTED = RGBColor(103, 103, 114)


def _set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def _hide_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'nil')
        borders.append(element)
    tblPr.append(borders)


def _font(run, *, name='Inter', size=10, bold=False, color=TEXT):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    return run


def _paragraph(container, text='', *, size=10, bold=False, color=TEXT, space_after=2, font='Inter'):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    _font(p.add_run(text), name=font, size=size, bold=bold, color=color)
    return p


def _bullet(container, text: str):
    p = _paragraph(container, '', size=9.5, space_after=1)
    _font(p.add_run('• '), size=9.5, bold=True)
    _font(p.add_run(text), size=9.5)
    return p


def _heading(container, text: str, size=17):
    p = _paragraph(container, text, size=size, bold=True, space_after=6, font='Montserrat')
    return p


def _section_rule(container):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), ACCENT)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _project_table(document: Document, project: ProjectItem):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Mm(55)
    table.columns[1].width = Mm(125)
    _hide_table_borders(table)
    left, right = table.rows[0].cells
    left.width = Mm(55)
    right.width = Mm(125)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(left, 70, 0, 80, 280)
    _set_cell_margins(right, 70, 280, 80, 0)

    if project.role:
        _paragraph(left, project.role, size=10, bold=True, space_after=3)
    if project.dates:
        _paragraph(left, project.dates, size=9.5, color=MUTED, space_after=7)
    if project.team:
        _paragraph(left, project.team, size=9, space_after=7)
    if project.technologies:
        _paragraph(left, ', '.join(project.technologies), size=8.8, color=MUTED, space_after=3)

    if project.name:
        _paragraph(right, project.name, size=11, bold=True, space_after=5)
    if project.description:
        _paragraph(right, 'Описание проекта', size=9.5, bold=True, space_after=2)
        _paragraph(right, project.description, size=9.5, space_after=6)
    if project.responsibilities:
        _paragraph(right, 'Выполняемые задачи (что и как делалось)', size=9.5, bold=True, space_after=3)
        for item in project.responsibilities:
            _bullet(right, item)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def render_docx(cv: CanonicalCv) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)

    normal = document.styles['Normal']
    normal.font.name = 'Inter'
    normal.font.size = Pt(10)
    normal.font.color.rgb = TEXT

    brand = _paragraph(document, 'IRLIX', size=12, bold=True, color=RGBColor(0, 142, 140), space_after=16, font='Montserrat')
    brand.alignment = 2
    _paragraph(document, cv.full_name or 'CV', size=34, bold=True, space_after=2, font='Montserrat')
    if cv.target_role:
        _paragraph(document, cv.target_role, size=14, color=MUTED, space_after=12, font='Montserrat')

    if cv.experience.commercial:
        _paragraph(document, f'Опыт коммерческой разработки: {cv.experience.commercial}', size=10, space_after=1)
    if cv.experience.role:
        label = cv.target_role or 'Опыт по роли'
        _paragraph(document, f'{label}: {cv.experience.role}', size=10, space_after=10)

    if cv.skill_groups or cv.tools:
        _heading(document, 'Навыки и умения:')
        _section_rule(document)
        for group in cv.skill_groups:
            _paragraph(document, group.title, size=10, bold=True, space_after=2)
            for item in group.items:
                _bullet(document, item)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        if cv.tools:
            _paragraph(document, 'Используемые инструменты:', size=10, bold=True, space_after=2)
            for tool in cv.tools:
                _bullet(document, tool)

    if cv.education:
        _heading(document, 'Образование:')
        _section_rule(document)
        for item in cv.education:
            title = item.institution or item.specialty or 'Образование'
            _paragraph(document, title, size=10, bold=True, space_after=1)
            details = [value for value in [item.degree, item.specialty, item.status] if value]
            for detail in details + item.details:
                _bullet(document, detail)

    if cv.languages:
        _paragraph(document, 'Иностранные языки:', size=10, bold=True, space_after=2)
        for item in cv.languages:
            value = f'{item.language} — {item.level}' if item.level else item.language
            _bullet(document, value)

    if cv.projects:
        document.add_page_break()
        _heading(document, f'Опыт работы — {cv.experience.commercial or ""}'.strip(' —'))
        _section_rule(document)
        for project in cv.projects:
            _project_table(document, project)

    if cv.additional or cv.extra_sections:
        document.add_page_break()
        _heading(document, 'Дополнительная информация')
        _section_rule(document)
        for item in cv.additional:
            if item.question:
                _paragraph(document, item.question, size=10, bold=True, space_after=2)
            _paragraph(document, item.answer, size=10, space_after=7)
        for section_item in cv.extra_sections:
            _paragraph(document, section_item.title, size=10, bold=True, space_after=2)
            for value in section_item.items:
                _bullet(document, value)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf(cv: CanonicalCv) -> bytes:
    docx_bytes = render_docx(cv)
    with tempfile.TemporaryDirectory(prefix='irlix-cv-') as tmp:
        tmp_path = Path(tmp)
        docx_path = tmp_path / 'cv.docx'
        docx_path.write_bytes(docx_bytes)
        soffice = shutil.which('libreoffice') or shutil.which('soffice')
        if not soffice:
            raise RuntimeError('LibreOffice is not installed in cv-converter service')
        subprocess.run(
            [soffice, '--headless', '--convert-to', 'pdf', '--outdir', str(tmp_path), str(docx_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=60,
        )
        pdf_path = tmp_path / 'cv.pdf'
        if not pdf_path.exists():
            raise RuntimeError('LibreOffice did not create PDF')
        return pdf_path.read_bytes()
