from __future__ import annotations

import re
from io import BytesIO

import fitz
from docx import Document


class UnsupportedSourceError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith('.pdf'):
        return _extract_pdf(data)
    if name.endswith('.docx'):
        return _extract_docx(data)
    if name.endswith('.doc'):
        raise UnsupportedSourceError('Legacy .doc пока не поддерживается. Сохраните документ как DOCX.')
    raise UnsupportedSourceError('Поддерживаются PDF и DOCX.')


def _clean_line(value: str) -> str:
    value = value.replace('\ufffe', '-').replace('\u00ad', '')
    return re.sub(r'[ \t]+', ' ', value).strip()


def _extract_pdf(data: bytes) -> str:
    document = fitz.open(stream=data, filetype='pdf')
    page_blocks: list[str] = []
    for page in document:
        blocks: list[str] = []
        for block in page.get_text('blocks', sort=True):
            lines = [_clean_line(line) for line in block[4].splitlines()]
            lines = [line for line in lines if line]
            if lines:
                blocks.append('\n'.join(lines))
        if blocks:
            page_blocks.append('\n\n'.join(blocks))
    text = '\n\n'.join(page_blocks).strip()
    if not text:
        raise UnsupportedSourceError('В PDF не найден текстовый слой. Для этого документа потребуется OCR.')
    return text


def _extract_docx(data: bytes) -> str:
    document = Document(BytesIO(data))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        value = _clean_line(paragraph.text)
        if value:
            blocks.append(value)
    for table in document.tables:
        for row in table.rows:
            cells = [_clean_line(cell.text) for cell in row.cells if _clean_line(cell.text)]
            if cells:
                blocks.append('\n'.join(cells))
    text = '\n\n'.join(blocks).strip()
    if not text:
        raise UnsupportedSourceError('DOCX не содержит извлекаемого текста.')
    return text
